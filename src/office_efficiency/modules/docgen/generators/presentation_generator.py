"""PPT 발표자료(.pptx) 자동 생성기."""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from office_efficiency.config import OUTPUT_DIR, TEMPLATES_DIR
from office_efficiency.core.logger import get_logger
from office_efficiency.modules.docgen.models import PresentationData

logger = get_logger("docgen.presentation")


def generate_presentation(
    data: PresentationData,
    template_path: Path | None = None,
    output_filename: str | None = None,
) -> Path:
    """발표자료 데이터를 기반으로 .pptx 프레젠테이션을 생성합니다."""
    default_template = TEMPLATES_DIR / "ppt" / "presentation_template.pptx"
    target_template = (
        template_path
        if (template_path and template_path.exists())
        else (default_template if default_template.exists() else None)
    )

    if target_template:
        prs = Presentation(target_template)
        logger.info(f"기존 PPT 서식 템플릿 로드: {target_template}")
    else:
        prs = Presentation()
        # 16:9 와이드스크린 비율 설정 (기본값)
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

    # 1. 표지 슬라이드 생성 (Title Slide)
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)

    title_shape = slide.shapes.title
    title_shape.text = data.title

    if len(slide.placeholders) > 1:
        subtitle_shape = slide.placeholders[1]
        sub_text = []
        if data.subtitle:
            sub_text.append(data.subtitle)
        meta = []
        if data.presenter:
            meta.append(f"발표자: {data.presenter}")
        if data.date:
            meta.append(f"일시: {data.date}")
        if meta:
            sub_text.append(" | ".join(meta))
        subtitle_shape.text = "\n".join(sub_text)

    # 2. 본문 슬라이드 생성
    content_layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]

    for slide_data in data.slides:
        slide = prs.slides.add_slide(content_layout)

        # 제목 설정
        slide.shapes.title.text = slide_data.title

        # 본문 불릿 포인트 설정
        if len(slide.placeholders) > 1:
            body_shape = slide.placeholders[1]
            tf = body_shape.text_frame
            tf.clear()

            for i, point in enumerate(slide_data.points):
                p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
                p.text = point
                p.font.size = Pt(20)
                p.level = 0

        # 발표자 노트가 있는 경우 추가
        if slide_data.notes:
            notes_slide = slide.notes_slide
            notes_text_frame = notes_slide.notes_text_frame
            notes_text_frame.text = slide_data.notes

    # 파일명 결정 및 저장
    safe_title = data.title.replace(" ", "_")
    filename = output_filename or f"발표자료_{safe_title}.pptx"
    output_path = OUTPUT_DIR / filename

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    logger.info(f"[green]PPT 발표자료 생성 완료: {output_path}[/green]")
    return output_path
