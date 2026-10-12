"""계약서 문서(.docx) 자동 생성기."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docxtpl import DocxTemplate

from office_efficiency.config import OUTPUT_DIR, TEMPLATES_DIR
from office_efficiency.core.logger import get_logger
from office_efficiency.modules.docgen.models import ContractData

logger = get_logger("docgen.contract")


def create_default_contract_template(template_path: Path) -> None:
    """기본 표준 계약서 워드 템플릿(docxtpl 태그 포함)을 생성합니다."""
    template_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # 제목
    title = doc.add_heading("{{ title }}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_intro = doc.add_paragraph()
    p_intro.add_run("{{ party_a.company_name }} (이하 '갑'이라 칭함)와 ")
    p_intro.add_run("{{ party_b.company_name }} (이하 '을'이라 칭함)는 ")
    p_intro.add_run("상호 신의성실의 원칙에 입각하여 다음과 같이 계약을 체결한다.")

    doc.add_heading("제 1 조 (목적)", level=1)
    doc.add_paragraph("본 계약은 {{ purpose }}을(를) 목적으로 한다.")

    doc.add_heading("제 2 조 (계약 기간)", level=1)
    doc.add_paragraph("계약 기간은 {{ start_date }}부터 {{ end_date }}까지로 한다.")

    doc.add_heading("제 3 조 (계약 금액 및 지급 조건)", level=1)
    doc.add_paragraph("1. 총 계약 금액: {{ total_amount }}")
    doc.add_paragraph("2. 지급 조건: {{ payment_terms }}")

    doc.add_heading("제 4 조 (특약 사항)", level=1)
    doc.add_paragraph("{% for article in articles %}")
    doc.add_paragraph("• {{ article }}")
    doc.add_paragraph("{% endfor %}")

    p_date = doc.add_paragraph("\n계약 체결일자: {{ contract_date }}")
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 당사자 서명란 표
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"

    table.cell(0, 0).text = "[갑] 발주처"
    table.cell(0, 1).text = "[을] 수급처"

    table.cell(1, 0).text = "상호: {{ party_a.company_name }}"
    table.cell(1, 1).text = "상호: {{ party_b.company_name }}"

    table.cell(2, 0).text = "대표자: {{ party_a.representative }} (인)"
    table.cell(2, 1).text = "대표자: {{ party_b.representative }} (인)"

    table.cell(3, 0).text = "사업자번호: {{ party_a.business_no }}"
    table.cell(3, 1).text = "사업자번호: {{ party_b.business_no }}"

    table.cell(4, 0).text = "주소: {{ party_a.address }}"
    table.cell(4, 1).text = "주소: {{ party_b.address }}"

    doc.save(template_path)
    logger.info(f"기본 계약서 템플릿 생성 완료: {template_path}")


def generate_contract(
    data: ContractData,
    template_path: Path | None = None,
    output_filename: str | None = None,
) -> Path:
    """계약서 데이터를 기반으로 .docx 문서를 생성합니다."""
    default_template = TEMPLATES_DIR / "contract" / "contract_template.docx"
    target_template = template_path or default_template

    if not target_template.exists():
        create_default_contract_template(target_template)

    doc = DocxTemplate(target_template)
    context = data.model_dump()
    doc.render(context)

    safe_title = data.title.replace(" ", "_")
    safe_no = data.contract_no.replace("/", "-")
    filename = output_filename or f"계약서_{safe_no}_{safe_title}.docx"
    output_path = OUTPUT_DIR / filename

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    logger.info(f"[green]계약서 생성 완료: {output_path}[/green]")
    return output_path
