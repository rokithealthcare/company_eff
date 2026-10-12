"""회의록 문서(.docx) 자동 생성기."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docxtpl import DocxTemplate

from office_efficiency.config import OUTPUT_DIR, TEMPLATES_DIR
from office_efficiency.core.logger import get_logger
from office_efficiency.modules.docgen.models import MeetingData

logger = get_logger("docgen.meeting")


def create_default_meeting_template(template_path: Path) -> None:
    """기본 회의록 워드 템플릿(docxtpl 태그 포함)을 생성합니다."""
    template_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # 제목
    title = doc.add_heading("회의록 (Meeting Minutes)", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 기본 정보 표
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"

    table.cell(0, 0).text = "회의명"
    table.cell(0, 1).text = "{{ title }}"
    table.cell(1, 0).text = "일시 / 장소"
    table.cell(1, 1).text = "{{ date }} / {{ location }}"
    table.cell(2, 0).text = "참석자"
    table.cell(2, 1).text = "{{ attendees_str }}"

    doc.add_heading("1. 회의 안건", level=1)
    doc.add_paragraph("{% for item in agenda %}")
    doc.add_paragraph("• {{ item }}")
    doc.add_paragraph("{% endfor %}")

    doc.add_heading("2. 주요 논의 내용", level=1)
    doc.add_paragraph("{% for item in discussions %}")
    doc.add_paragraph("• {{ item }}")
    doc.add_paragraph("{% endfor %}")

    doc.add_heading("3. 결정 사항", level=1)
    doc.add_paragraph("{% for item in decisions %}")
    doc.add_paragraph("• {{ item }}")
    doc.add_paragraph("{% endfor %}")

    doc.add_heading("4. 향후 실행 과제 (Action Items)", level=1)
    action_table = doc.add_table(rows=2, cols=3)
    action_table.style = "Table Grid"
    action_table.cell(0, 0).text = "과제명"
    action_table.cell(0, 1).text = "담당자"
    action_table.cell(0, 2).text = "완료목표일"

    action_table.cell(1, 0).text = "{% for action in action_items %}{{ action.task }}"
    action_table.cell(1, 1).text = "{{ action.assignee }}"
    action_table.cell(
        1, 2
    ).text = "{{ action.due_date }}{% if not loop.last %}\n{% endif %}{% endfor %}"

    doc.save(template_path)
    logger.info(f"기본 회의록 템플릿 생성 완료: {template_path}")


def generate_meeting_minutes(
    data: MeetingData,
    template_path: Path | None = None,
    output_filename: str | None = None,
) -> Path:
    """회의록 데이터를 기반으로 .docx 문서를 생성합니다."""
    default_template = TEMPLATES_DIR / "meeting" / "meeting_template.docx"
    target_template = template_path or default_template

    # 템플릿 파일이 없을 경우 기본 서식 생성
    if not target_template.exists():
        create_default_meeting_template(target_template)

    # 템플릿 렌더링
    doc = DocxTemplate(target_template)
    context = data.model_dump()
    context["attendees_str"] = ", ".join(data.attendees) if data.attendees else "참석자 미지정"

    doc.render(context)

    # 파일명 생성 및 저장
    safe_title = data.title.replace(" ", "_")
    safe_date = data.date.replace("-", "")
    filename = output_filename or f"회의록_{safe_title}_{safe_date}.docx"
    output_path = OUTPUT_DIR / filename

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    logger.info(f"[green]회의록 생성 완료: {output_path}[/green]")
    return output_path
