"""구매사양서 문서(.docx) 자동 생성기."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docxtpl import DocxTemplate

from office_efficiency.config import OUTPUT_DIR, TEMPLATES_DIR
from office_efficiency.core.logger import get_logger
from office_efficiency.modules.docgen.models import PurchaseSpecData

logger = get_logger("docgen.spec")


def create_default_spec_template(template_path: Path) -> None:
    """기본 구매사양서 워드 템플릿(docxtpl 태그 포함)을 생성합니다."""
    template_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # 제목
    title = doc.add_heading("구 매 사 양 서", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 기본 메타데이터 표
    table = doc.add_table(rows=4, cols=4)
    table.style = "Table Grid"

    table.cell(0, 0).text = "문서번호"
    table.cell(0, 1).text = "{{ doc_no }}"
    table.cell(0, 2).text = "요청일자"
    table.cell(0, 3).text = "{{ request_date }}"

    table.cell(1, 0).text = "요청부서"
    table.cell(1, 1).text = "{{ department }}"
    table.cell(1, 2).text = "요청자"
    table.cell(1, 3).text = "{{ requester }}"

    table.cell(2, 0).text = "구매건명"
    table.cell(2, 1).text = "{{ title }}"
    table.cell(2, 2).text = "희망납기"
    table.cell(2, 3).text = "{{ delivery_date }}"

    table.cell(3, 0).text = "납품장소"
    table.cell(3, 1).text = "{{ delivery_place }}"
    table.cell(3, 2).text = "구매목적"
    table.cell(3, 3).text = "{{ purpose }}"

    doc.add_heading("1. 구매 대상 품목 및 세부 사양", level=1)

    item_table = doc.add_table(rows=2, cols=6)
    item_table.style = "Table Grid"
    item_table.cell(0, 0).text = "No"
    item_table.cell(0, 1).text = "품명"
    item_table.cell(0, 2).text = "규격 및 상세 사양"
    item_table.cell(0, 3).text = "수량"
    item_table.cell(0, 4).text = "단위"
    item_table.cell(0, 5).text = "비고"

    item_table.cell(1, 0).text = "{% for item in items %}{{ loop.index }}"
    item_table.cell(1, 1).text = "{{ item.name }}"
    item_table.cell(1, 2).text = "{{ item.spec }}"
    item_table.cell(1, 3).text = "{{ item.quantity }}"
    item_table.cell(1, 4).text = "{{ item.unit }}"
    item_table.cell(1, 5).text = "{{ item.remarks }}{% if not loop.last %}\n{% endif %}{% endfor %}"

    doc.add_heading("2. 특기 및 검수 조건", level=1)
    doc.add_paragraph("{% for req in special_requirements %}")
    doc.add_paragraph("• {{ req }}")
    doc.add_paragraph("{% endfor %}")

    doc.save(template_path)
    logger.info(f"기본 구매사양서 템플릿 생성 완료: {template_path}")


def generate_purchase_spec(
    data: PurchaseSpecData,
    template_path: Path | None = None,
    output_filename: str | None = None,
) -> Path:
    """구매사양서 데이터를 기반으로 .docx 문서를 생성합니다."""
    default_template = TEMPLATES_DIR / "spec" / "spec_template.docx"
    target_template = template_path or default_template

    if not target_template.exists():
        create_default_spec_template(target_template)

    doc = DocxTemplate(target_template)
    context = data.model_dump()
    doc.render(context)

    safe_title = data.title.replace(" ", "_")
    safe_doc_no = data.doc_no.replace("/", "-")
    filename = output_filename or f"구매사양서_{safe_doc_no}_{safe_title}.docx"
    output_path = OUTPUT_DIR / filename

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    logger.info(f"[green]구매사양서 생성 완료: {output_path}[/green]")
    return output_path
